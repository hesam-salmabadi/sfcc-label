"""Standardize temperatures and CS616 moisture from Chelene's four plot books.

Each plot book carries its own temperature and moisture streams. Its unlabeled
clock is interpreted as fixed Ontario standard time (UTC-5), based on the
continuous daylight-saving transition hours and summer soil-temperature cycle.
Distinct probes remain separate; no replicate averaging or same-depth
temperature/moisture pairing is performed.
"""

from __future__ import annotations

import csv
import math
import re
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from .io import OBSERVATION_COLUMNS, write_metadata
from .models import SensorMetadata
from .naming import sensor_id as make_sensor_id


PLOTS = {"AS3": "Aw", "BS1": "Sb", "JP1": "Pj", "MW1": "Mw"}
TEMP_DEPTHS = {6, 15, 30}
MOISTURE_DEPTHS = {10, 18}
SOURCE_OFFSET = timedelta(hours=-5)
CLOCK_EVIDENCE = (
    "inferred fixed Ontario standard time: all four workbooks retain 02:00 on "
    "spring DST dates and have only one 01:00 on fall DST dates; 2018-2021 "
    "summer 6 cm temperature peaks at source hours 15-19 and troughs at 7-9"
)
CONTEXT_COLUMNS = ("sensor_id", "measurement", "plot", "subplot", "replicate",
                   "depth_basis", "installation_depth_cm", "support_from_cm",
                   "support_to_cm", "midpoint_cm", "source_column", "raw_column",
                   "source_file", "nearby_temperature_sensor_id",
                   "aspen_coordinate_note", "clock_basis", "clock_evidence",
                   "undated_rows_skipped")
STATUS_COLUMNS = ("sensor_id", "status", "source_rows", "observed_hours", "output_hours",
                  "temperature_hours", "moisture_hours", "raw_hours", "first_utc",
                  "last_utc", "detail")


def _workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Chapleau importer requires pip install 'sfcc-label[source-excel]'") from exc
    return load_workbook(path, read_only=True, data_only=True)


def _rounded_hour(stamp: datetime) -> datetime:
    if stamp.tzinfo is not None:
        raise ValueError("expected a naive publisher timestamp")
    return (stamp + timedelta(minutes=30)).replace(minute=0, second=0, microsecond=0)


def _numeric(value, where: str) -> float | None:
    if value is None or isinstance(value, str) and value.strip().upper() in {"", "NA", "NAN"}:
        return None
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{where}: invalid numeric value {value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{where}: nonfinite numeric value")
    return result


def _local_to_utc(stamp: datetime) -> datetime:
    """Convert inferred fixed Ontario standard time (UTC-5) to UTC."""
    return (stamp - SOURCE_OFFSET).replace(tzinfo=timezone.utc)


def _put(buckets: dict, stamp: datetime, value, where: str) -> None:
    if stamp in buckets and buckets[stamp] != value:
        raise ValueError(f"{where}: conflicting rows for {stamp}")
    buckets[stamp] = value


def _write_hourly(path: Path, samples: dict[datetime, tuple[float | None, float | None, float | None]]):
    first, last = min(samples), max(samples)
    counts = [0, 0, 0]
    output_hours = 0
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(OBSERVATION_COLUMNS)
        stamp = first
        while stamp <= last:
            values = samples.get(stamp, (None, None, None))
            counts = [count + (value is not None) for count, value in zip(counts, values)]
            writer.writerow((stamp.strftime("%Y-%m-%dT%H:00:00Z"),
                             *("NaN" if value is None else format(value, ".15g") for value in values)))
            output_hours += 1
            stamp += timedelta(hours=1)
    return {"observed_hours": len(samples), "output_hours": output_hours,
            "temperature_hours": counts[0], "moisture_hours": counts[1],
            "raw_hours": counts[2],
            "first_utc": first.strftime("%Y-%m-%dT%H:00:00Z"),
            "last_utc": last.strftime("%Y-%m-%dT%H:00:00Z")}


def import_chapleau(chelene_dir: str | Path, observations_dir: str | Path, sensors_file: str | Path,
                    context_file: str | Path, status_file: str | Path,
                    aspen_latitude: float, aspen_longitude: float,
                    aspen_coordinate_evidence: str) -> list[dict]:
    """Import all 12 temperatures and 48 moisture probes from four Chelene books."""
    if not aspen_coordinate_evidence.strip():
        raise ValueError("provide evidence for selected Aspen coordinates")
    source_root = Path(chelene_dir)
    destinations = [Path(observations_dir), Path(sensors_file), Path(context_file), Path(status_file)]
    if any(path.exists() for path in destinations):
        raise FileExistsError("Chapleau output already exists; use new destinations")
    for destination in destinations:
        destination.parent.mkdir(parents=True, exist_ok=True)
    sensors, contexts, statuses = [], [], []
    with tempfile.TemporaryDirectory(prefix=".sfcc-chapleau-", dir=Path(observations_dir).parent) as temporary:
        stage = Path(temporary)
        for plot, suffix in PLOTS.items():
            paths = list(source_root.glob(f"Chapleau_{suffix}_soil.moisture.temp.xlsx"))
            if len(paths) != 1:
                raise ValueError(f"expected one Chelene workbook for {plot}, found {len(paths)}")
            path = paths[0]
            book = _workbook(path)
            try:
                readme = book.worksheets[0]
                location = readme.cell(3, 1).value
                match = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)\s*", str(location))
                if not match:
                    raise ValueError(f"{path}: missing plot coordinates in the first sheet")
                latitude, longitude = (aspen_latitude, aspen_longitude) if plot == "AS3" else (
                    float(match[1]), float(match[2]))
                data = book["Data"]
                header = next(data.values)
                columns = {name: i for i, name in enumerate(header) if name is not None}
                temperature_columns = {6: "p1_temp_6", 15: "p3_temp_15", 30: "p5_temp_30"}
                if any(name not in columns for name in temperature_columns.values()):
                    raise ValueError(f"{path}: missing a subplot temperature column")
                pattern = re.compile(r"p([135])([abcd])_vwc\.cal_(10|18)")
                probes = []
                for name in columns:
                    match = pattern.fullmatch(name)
                    if match:
                        subplot, replicate, extent = int(match[1]), match[2], int(match[3])
                        raw_name = f"p{subplot}{replicate}_period_{extent}"
                        if raw_name not in columns:
                            raise ValueError(f"{path}: {name} has no raw period")
                        probes.append((subplot, replicate, extent, name, raw_name))
                if len(probes) != 12:
                    raise ValueError(f"{path}: expected 12 moisture probes, found {len(probes)}")
                temperatures = {depth: {} for depth in TEMP_DEPTHS}
                by_probe = {name: {} for _, _, _, name, _ in probes}
                undated = 0
                dated_rows = 0
                for line, row in enumerate(data.iter_rows(min_row=2, values_only=True), 2):
                    local = row[0]
                    if not isinstance(local, datetime):
                        undated += 1
                        continue
                    dated_rows += 1
                    local = _rounded_hour(local)
                    utc = _local_to_utc(local)
                    stamp = utc.replace(tzinfo=None)
                    for depth, name in temperature_columns.items():
                        value = _numeric(row[columns[name]], f"{path}:{line}:{name}")
                        if value is not None:
                            _put(temperatures[depth], stamp, (value, None, None),
                                 f"{path}:{line}:{name}")
                    for _, _, _, name, raw_name in probes:
                        vwc = _numeric(row[columns[name]], f"{path}:{line}:{name}")
                        period = _numeric(row[columns[raw_name]], f"{path}:{line}:{raw_name}")
                        _put(by_probe[name], stamp,
                             (None, None if vwc is None else vwc / 100, period),
                             f"{path}:{line}:{name}")
                for depth, name in temperature_columns.items():
                    sensor_id = make_sensor_id("chapleau", plot, float(depth), tag="t")
                    sensors.append(SensorMetadata(
                        sensor_id=sensor_id, source="chapleau", latitude=latitude,
                        longitude=longitude, site_id=plot, network="Chapleau", station=plot,
                        depth_cm=float(depth), depth_from_cm=float(depth), depth_to_cm=float(depth),
                        source_id=name, source_url=str(path),
                        timezone_original="UTC-5 fixed (inferred Ontario standard time)",
                        soil_temperature_sensor_type="soil temperature probe",
                        sensor_type_source="Chelene workbook ReadMe"))
                    counts = _write_hourly(stage / f"{sensor_id}.csv", temperatures[depth])
                    contexts.append({"sensor_id": sensor_id, "measurement": "temperature",
                                     "plot": plot, "subplot": {6: 1, 15: 3, 30: 5}[depth],
                                     "replicate": "NaN", "depth_basis": "point temperature depth, Chelene workbook",
                                     "installation_depth_cm": depth, "support_from_cm": depth,
                                     "support_to_cm": depth, "midpoint_cm": depth,
                                     "source_column": name, "raw_column": "NaN", "source_file": str(path),
                                     "nearby_temperature_sensor_id": "NaN",
                                     "aspen_coordinate_note": aspen_coordinate_evidence if plot == "AS3" else "",
                                     "clock_basis": "fixed UTC-5; add five hours",
                                     "clock_evidence": CLOCK_EVIDENCE,
                                     "undated_rows_skipped": undated})
                    statuses.append({"sensor_id": sensor_id, "status": "imported",
                                     "source_rows": dated_rows + undated, **counts,
                                     "detail": f"undated={undated};clock=inferred fixed UTC-5"})
                for subplot, replicate, extent, name, raw_name in probes:
                    midpoint = extent / 2
                    sensor_id = make_sensor_id("chapleau", plot, 0.0, float(extent), tag=f"p{subplot}{replicate}")
                    sensors.append(SensorMetadata(
                        sensor_id=sensor_id, source="chapleau", latitude=latitude,
                        longitude=longitude, site_id=plot, network="Chapleau", station=plot,
                        depth_cm=midpoint, depth_from_cm=0.0, depth_to_cm=float(extent),
                        raw_variable="CS616_period", raw_unit="microsecond",
                        source_id=name, source_url=str(path),
                        timezone_original="UTC-5 fixed (inferred Ontario standard time)",
                        soil_moisture_method="calibrated VWC percent / 100; source soil-profile coefficients",
                        soil_moisture_sensor_type="CS616",
                        sensor_type_source="Chapleau source workbook/soil-profile metadata",
                        sensor_type_note="Raw period output is the CS616 measurement; VWC is separately calibrated."))
                    values = by_probe[name]
                    counts = _write_hourly(stage / f"{sensor_id}.csv", values)
                    contexts.append({"sensor_id": sensor_id, "measurement": "moisture",
                                     "plot": plot, "subplot": subplot, "replicate": replicate,
                                     "depth_basis": "30 cm CS616 rod inserted diagonally from surface",
                                     "installation_depth_cm": extent, "support_from_cm": 0,
                                     "support_to_cm": extent, "midpoint_cm": midpoint,
                                     "source_column": name, "raw_column": raw_name,
                                     "source_file": str(path),
                                     "nearby_temperature_sensor_id": make_sensor_id(
                                         "chapleau", plot, float({1: 6, 3: 15, 5: 30}[subplot]), tag="t"),
                                     "aspen_coordinate_note": aspen_coordinate_evidence if plot == "AS3" else "",
                                     "clock_basis": "fixed UTC-5; add five hours",
                                     "clock_evidence": CLOCK_EVIDENCE,
                                     "undated_rows_skipped": undated})
                    statuses.append({"sensor_id": sensor_id, "status": "imported",
                                     "source_rows": dated_rows + undated, **counts,
                                     "detail": f"undated={undated};clock=inferred fixed UTC-5"})
            finally:
                book.close()
            print(f"Prepared Chapleau {plot} temperature and moisture", flush=True)
        if len(sensors) != 60 or len({sensor.sensor_id for sensor in sensors}) != 60:
            raise ValueError("expected 60 unique Chapleau streams")
        write_metadata(stage / "sensors.csv", sensors)
        for filename, columns, rows in (("context.csv", CONTEXT_COLUMNS, contexts),
                                        ("status.csv", STATUS_COLUMNS, statuses)):
            with (stage / filename).open("w", newline="", encoding="utf-8") as stream:
                writer = csv.DictWriter(stream, fieldnames=columns)
                writer.writeheader()
                writer.writerows(rows)
        output = Path(observations_dir)
        output.mkdir()
        for sensor in sensors:
            (stage / f"{sensor.sensor_id}.csv").replace(output / f"{sensor.sensor_id}.csv")
        (stage / "sensors.csv").replace(sensors_file)
        (stage / "context.csv").replace(context_file)
        (stage / "status.csv").replace(status_file)
    return statuses
