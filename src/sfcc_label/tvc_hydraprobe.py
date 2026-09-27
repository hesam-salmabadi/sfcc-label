"""Import the TVC Soil Stations HydraProbe workbooks.

Each workbook contains four probes. H1/H3 are 5 cm and H2/H4 are 10 cm.
The raw-value column is effective bulk permittivity from the paper's
complex-permittivity conversion, using the temperature-corrected real and
imaginary components:

    eps_eff = (eps_real + sqrt(eps_real**2 + eps_imag**2)) / 2

Workbook timestamps are explicitly UTC and are recorded every 30 minutes.
"""

from __future__ import annotations

import csv
import math
import tempfile
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from .io import OBSERVATION_COLUMNS, write_metadata
from .models import SensorMetadata
from .naming import sensor_id as make_sensor_id


COORDINATES = {
    "CalTarget": (68.7396682571399, -133.492117976203),
    "DriftSite": (68.7376129080591, -133.48768343498),
    "MainMet": (68.7465154492267, -133.502211706235),
    "OldTrench": (68.7307937401294, -133.509789777214),
    "SouthTundra": (68.7263651405424, -133.507071645923),
    "ValleyBottom": (68.7373228447467, -133.508587508996),
}
DEPTHS = {"H1": 5.0, "H2": 10.0, "H3": 5.0, "H4": 10.0}
CONTEXT_COLUMNS = (
    "sensor_id", "station", "probe", "depth_cm", "source_file",
    "source_rows", "observed_hours", "temperature_hours", "moisture_hours",
    "permittivity_hours", "first_utc", "last_utc", "time_basis",
    "permittivity_formula",
)
FORMULA = "(epsilon_real_corrected + sqrt(epsilon_real_corrected^2 + epsilon_imag_corrected^2)) / 2"


def _number(value: object) -> float | None:
    if value is None or str(value).strip().lower() in {"", "nan", "na", "null"}:
        return None
    result = float(value)
    if not math.isfinite(result):
        return None
    return result


def _station_from_file(path: Path) -> str:
    for station in COORDINATES:
        if path.name.startswith(station + "_"):
            return station
    raise ValueError(f"unrecognized TVC station workbook: {path.name}")


def _read_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("TVC HydraProbe import requires openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    first = next(rows, None)
    header = next(rows, None)
    units = next(rows, None)
    next(rows, None)
    positions = next(rows, None)
    if not header or header[0] != "TIMESTAMP" or not positions:
        raise ValueError(f"unexpected TVC HydraProbe header in {path}")
    needed = ["TIMESTAMP"]
    for probe in DEPTHS:
        needed.extend((f"{probe}_Soil_Moisture", f"{probe}_Soil_Temperature_C",
                       f"{probe}_Cor_Real_Permittivity",
                       f"{probe}_Cor_Imaginary_Permittivity"))
    indices = {name: header.index(name) for name in needed if name in header}
    missing = set(needed) - set(indices)
    if missing:
        raise ValueError(f"{path}: missing columns {sorted(missing)}")
    result = []
    for row_number, row in enumerate(rows, start=6):
        if not row or row[indices["TIMESTAMP"]] is None:
            continue
        stamp = row[indices["TIMESTAMP"]]
        if not isinstance(stamp, datetime):
            raise ValueError(f"{path}:{row_number}: TIMESTAMP is not a datetime")
        if stamp.tzinfo is not None:
            stamp = stamp.replace(tzinfo=None)
        record = {"timestamp": stamp, "source_row": row_number}
        for probe in DEPTHS:
            real = _number(row[indices[f"{probe}_Cor_Real_Permittivity"]])
            imag = _number(row[indices[f"{probe}_Cor_Imaginary_Permittivity"]])
            permittivity = None if real is None or imag is None else (real + math.hypot(real, imag)) / 2.0
            record[probe] = (_number(row[indices[f"{probe}_Soil_Temperature_C"]]),
                             _number(row[indices[f"{probe}_Soil_Moisture"]]), permittivity)
        result.append(record)
    workbook.close()
    if not result:
        raise ValueError(f"no observations in {path}")
    return result


def _write_hourly(path: Path, values: dict[datetime, list[list[float]]]) -> None:
    first, last = min(values), max(values)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(OBSERVATION_COLUMNS)
        stamp = first
        while stamp <= last:
            buckets = values.get(stamp, [[], [], []])
            means = [math.fsum(bucket) / len(bucket) if bucket else None for bucket in buckets]
            writer.writerow((stamp.strftime("%Y-%m-%dT%H:00:00Z"),
                             *("NaN" if value is None else format(value, ".15g") for value in means)))
            stamp += timedelta(hours=1)


def import_tvc_hydraprobe(source_dir: str | Path, observations_dir: str | Path,
                          sensors_file: str | Path, context_file: str | Path) -> list[dict]:
    source = Path(source_dir)
    output = Path(observations_dir)
    destinations = (output, Path(sensors_file), Path(context_file))
    if any(path.exists() for path in destinations):
        raise FileExistsError("TVC HydraProbe output already exists")
    files = sorted(path for path in source.glob("*.xlsx") if not path.name.startswith(".~$"))
    if len(files) != len(COORDINATES):
        raise ValueError(f"expected {len(COORDINATES)} station workbooks; found {len(files)}")
    streams: dict[str, dict[datetime, list[list[float]]]] = defaultdict(lambda: defaultdict(lambda: [[], [], []]))
    source_rows = defaultdict(int)
    source_files = {}
    for file in files:
        station = _station_from_file(file)
        records = _read_workbook(file)
        for record in records:
            hour = record["timestamp"].replace(minute=0, second=0, microsecond=0)
            for probe in DEPTHS:
                key = (station, probe)
                source_files[key] = str(file)
                source_rows[key] += 1
                buckets = streams[key][hour]
                for index, value in enumerate(record[probe]):
                    if value is not None:
                        buckets[index].append(value)
    output.parent.mkdir(parents=True, exist_ok=True)
    Path(sensors_file).parent.mkdir(parents=True, exist_ok=True)
    Path(context_file).parent.mkdir(parents=True, exist_ok=True)
    sensors, contexts = [], []
    with tempfile.TemporaryDirectory(prefix=".sfcc-tvc-hydraprobe-", dir=output.parent) as temp:
        stage = Path(temp)
        for station, probe in sorted(streams):
            depth = DEPTHS[probe]
            sensor_id = make_sensor_id("tvc_hydraprobe", station, depth, tag=probe)
            values = streams[(station, probe)]
            latitude, longitude = COORDINATES[station]
            sensors.append(SensorMetadata(
                sensor_id=sensor_id, source="tvc_hydraprobe", latitude=latitude,
                longitude=longitude, site_id=station, network="TVC Soil Stations",
                station=station, depth_cm=depth, depth_from_cm=depth, depth_to_cm=depth,
                raw_variable="bulk_permittivity", raw_unit="relative permittivity",
                source_id=f"{station}_{probe}", source_url="https://github.com/hesam-salmabadi/TVC-Soil-Stations",
                timezone_original="UTC (publisher metadata workbook)",
                soil_moisture_method="HydraProbe volumetric water fraction",
                soil_temperature_sensor_type="HydraProbe",
                soil_moisture_sensor_type="HydraProbe",
                sensor_type_source="TVC Soil Stations metadata workbook",
                sensor_type_note="H1/H3 support 5 cm; H2/H4 support 10 cm.",
            ))
            _write_hourly(stage / f"{sensor_id}.csv", values)
            contexts.append({
                "sensor_id": sensor_id, "station": station, "probe": probe,
                "depth_cm": depth, "source_file": source_files[(station, probe)],
                "source_rows": source_rows[(station, probe)], "observed_hours": len(values),
                "temperature_hours": sum(bool(x[0]) for x in values.values()),
                "moisture_hours": sum(bool(x[1]) for x in values.values()),
                "permittivity_hours": sum(bool(x[2]) for x in values.values()),
                "first_utc": min(values).isoformat() + "Z", "last_utc": max(values).isoformat() + "Z",
                "time_basis": "TVC metadata workbook TIMESTAMP; UTC; 30-minute source averaged to UTC hour",
                "permittivity_formula": FORMULA,
            })
        write_metadata(stage / "sensors.csv", sensors)
        with (stage / "context.csv").open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=CONTEXT_COLUMNS)
            writer.writeheader(); writer.writerows(contexts)
        output.mkdir()
        for sensor in sensors:
            (stage / f"{sensor.sensor_id}.csv").replace(output / f"{sensor.sensor_id}.csv")
        (stage / "sensors.csv").replace(sensors_file)
        (stage / "context.csv").replace(context_file)
    return contexts
