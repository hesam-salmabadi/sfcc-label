"""Import the three original Dryden temperature series without inferring depth.

The workbook labels its streams 6, 18 and 30 cm. The data owner confirmed
these are temperature depths; the CFS field document's 10, 14 and 18 cm are
moisture-probe installation depths. Callers still supply an explicit depth
mapping for an auditable import. The workbook declares its timestamps UTC.
"""

from __future__ import annotations

import csv
import math
import tempfile
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from .io import OBSERVATION_COLUMNS, write_metadata
from .models import SensorMetadata
from .naming import sensor_id as make_sensor_id


SOURCE_DEPTHS = (6, 18, 30)
CONTEXT_COLUMNS = ("sensor_id", "source_id", "workbook_depth_cm", "assigned_depth_cm",
                   "depth_evidence", "source_file", "timezone_evidence",
                   "undated_source_rows")
STATUS_COLUMNS = ("sensor_id", "source_id", "status", "source_rows", "dated_rows",
                  "undated_rows", "observed_hours", "output_hours", "first_utc", "last_utc")


def _load_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Dryden importer requires pip install 'sfcc-label[source-excel]'") from exc
    return load_workbook(path, read_only=True, data_only=True)


def _metadata(book) -> dict[str, dict]:
    expected = {f"Dryden_{depth}" for depth in SOURCE_DEPTHS}
    if "Metadata" not in book.sheetnames or "Raw time series data" not in book.sheetnames:
        raise ValueError("Dryden workbook is missing Metadata or Raw time series data")
    rows = book["Metadata"].values
    header = next(rows)
    needed = {"Raw_data_identifier", "Latitude", "Longitude", "Sensor_height", "Timezone",
              "Unit", "Microclimate_measurement"}
    if not needed.issubset(header):
        raise ValueError(f"Dryden metadata is missing {sorted(needed - set(header))}")
    metadata = {}
    for row in rows:
        record = dict(zip(header, row))
        ident = record["Raw_data_identifier"]
        if ident not in expected or ident in metadata:
            raise ValueError(f"unexpected or repeated Dryden identifier: {ident!r}")
        depth = int(ident.rsplit("_", 1)[1])
        if record["Sensor_height"] != -depth:
            raise ValueError(f"{ident}: workbook sensor height disagrees with its identifier")
        if record["Timezone"] != "UTC" or record["Unit"] != "°C":
            raise ValueError(f"{ident}: expected workbook UTC and °C metadata")
        if record["Microclimate_measurement"] != "Temperature":
            raise ValueError(f"{ident}: not a temperature stream")
        latitude, longitude = float(record["Latitude"]), float(record["Longitude"])
        if not 0 <= latitude <= 90 or not -180 <= longitude <= 180:
            raise ValueError(f"{ident}: invalid northern coordinates")
        record["Latitude"], record["Longitude"] = latitude, longitude
        metadata[ident] = record
    if set(metadata) != expected:
        raise ValueError(f"expected {sorted(expected)}, found {sorted(metadata)}")
    if len({(row["Latitude"], row["Longitude"]) for row in metadata.values()}) != 1:
        raise ValueError("Dryden streams disagree on site coordinates")
    return metadata


def _read_samples(book) -> tuple[dict[str, dict[datetime, float]], dict[str, int], dict[str, int]]:
    expected = {f"Dryden_{depth}" for depth in SOURCE_DEPTHS}
    rows = book["Raw time series data"].values
    if next(rows)[:6] != ("Raw_data_identifier", "Year", "Month", "Day", "Time (24h)",
                          "Temperature"):
        raise ValueError("unexpected Dryden time-series columns")
    samples: dict[str, dict[datetime, float]] = defaultdict(dict)
    counts = defaultdict(int)
    undated = defaultdict(int)
    for line, row in enumerate(rows, 2):
        ident, year, month, day, clock, temperature = row[:6]
        if ident not in expected:
            raise ValueError(f"Dryden row {line}: unexpected identifier {ident!r}")
        counts[ident] += 1
        if (year, month, day, clock) == ("NA", "NA", "NA", "NA"):
            undated[ident] += 1
            continue
        try:
            stamp = datetime.strptime(f"{year}-{month}-{day} {clock}", "%Y-%m-%d %H:%M:%S")
            value = float(temperature)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Dryden row {line}: invalid dated observation") from exc
        if stamp.minute or stamp.second or not math.isfinite(value):
            raise ValueError(f"Dryden row {line}: non-hourly or nonfinite observation")
        if stamp in samples[ident] and samples[ident][stamp] != value:
            raise ValueError(f"Dryden row {line}: conflicting repeated hour for {ident}")
        samples[ident][stamp] = value
    if set(samples) != expected:
        raise ValueError("one or more Dryden streams have no dated observations")
    return samples, counts, undated


def import_dryden(workbook: str | Path, observations_dir: str | Path,
                  sensors_file: str | Path, context_file: str | Path,
                  status_file: str | Path, depth_map: dict[int, float],
                  depth_evidence: str) -> list[dict]:
    """Write three temperature-only UTC files after explicit depth resolution."""
    if set(depth_map) != set(SOURCE_DEPTHS) or not depth_evidence.strip():
        raise ValueError("provide assigned depths for source 6, 18 and 30, with depth evidence")
    if any(not math.isfinite(depth) or depth < 0 for depth in depth_map.values()):
        raise ValueError("assigned depths must be nonnegative finite centimetres")
    source = Path(workbook)
    destinations = [Path(observations_dir), Path(sensors_file), Path(context_file),
                    Path(status_file)]
    if any(path.exists() for path in destinations):
        raise FileExistsError("Dryden output already exists; use new destinations")
    book = _load_workbook(source)
    try:
        metadata = _metadata(book)
        samples, counts, undated = _read_samples(book)
    finally:
        book.close()
    for destination in destinations:
        destination.parent.mkdir(parents=True, exist_ok=True)
    sensors = []
    contexts = []
    statuses = []
    with tempfile.TemporaryDirectory(prefix=".sfcc-dryden-", dir=Path(observations_dir).parent) as temporary:
        stage = Path(temporary)
        for source_depth in SOURCE_DEPTHS:
            ident = f"Dryden_{source_depth}"
            assigned = float(depth_map[source_depth])
            sensor_id = make_sensor_id("dryden", "Dryden", assigned)
            record = metadata[ident]
            sensors.append(SensorMetadata(
                sensor_id=sensor_id, source="dryden", latitude=record["Latitude"],
                longitude=record["Longitude"], site_id="Dryden", network="Dryden",
                station="Dryden", depth_cm=assigned, depth_from_cm=assigned,
                depth_to_cm=assigned, source_id=ident, source_url=str(source),
                timezone_original="UTC (workbook metadata)",
            ))
            hours = samples[ident]
            first, last = min(hours), max(hours)
            output_hours = 0
            with (stage / f"{sensor_id}.csv").open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(OBSERVATION_COLUMNS)
                stamp = first
                while stamp <= last:
                    value = hours.get(stamp)
                    writer.writerow((stamp.strftime("%Y-%m-%dT%H:00:00Z"),
                                     "NaN" if value is None else format(value, ".15g"),
                                     "NaN", "NaN"))
                    output_hours += 1
                    stamp += timedelta(hours=1)
            contexts.append({
                "sensor_id": sensor_id, "source_id": ident,
                "workbook_depth_cm": source_depth, "assigned_depth_cm": assigned,
                "depth_evidence": depth_evidence, "source_file": str(source),
                "timezone_evidence": "Metadata/Timezone=UTC",
                "undated_source_rows": undated[ident],
            })
            statuses.append({
                "sensor_id": sensor_id, "source_id": ident, "status": "imported",
                "source_rows": counts[ident], "dated_rows": len(hours),
                "undated_rows": undated[ident], "observed_hours": len(hours),
                "output_hours": output_hours,
                "first_utc": first.strftime("%Y-%m-%dT%H:00:00Z"),
                "last_utc": last.strftime("%Y-%m-%dT%H:00:00Z"),
            })
        write_metadata(stage / "sensors.csv", sensors)
        for filename, columns, rows in (("context.csv", CONTEXT_COLUMNS, contexts),
                                        ("status.csv", STATUS_COLUMNS, statuses)):
            with (stage / filename).open("w", newline="", encoding="utf-8") as stream:
                writer = csv.DictWriter(stream, fieldnames=columns)
                writer.writeheader()
                writer.writerows(rows)
        output = Path(observations_dir)
        output.mkdir()
        for sensor in sensors:
            (stage / f"{sensor.sensor_id}.csv").replace(output / f"{sensor.sensor_id}.csv")
        (stage / "sensors.csv").replace(sensors_file)
        (stage / "context.csv").replace(context_file)
        (stage / "status.csv").replace(status_file)
    return statuses
